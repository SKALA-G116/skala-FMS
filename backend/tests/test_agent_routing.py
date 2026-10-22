"""
==============================================================================
FMS_FILE_SPEC
test_agent_routing.py - 기능·경계·회귀 검증
==============================================================================
설명:
    - 기능·경계·회귀 검증
입력·출력 명세:
    - 입력: 테스트 fixture·가상 입력·명시적 mock/테스트 대역
    - 출력: assert 결과와 실패 사유
실행 가정:
    - DB 통합 테스트는 별도 _test DB·임시 assets 사용; 대역과 실제 모델 검증을 구분
작성 정보:
    - 작성자: 김재현(Finn)
    - 이메일: penguin.klg@gmail.com
    - 작성일: 2026-10-21 (최초 Git 등록일 기준)
    - 라이선스: 미지정
    - 버전: 0.1.0
    - 상태: Development
변경 이력 (Revision History):
    DATE        AUTHOR      VERSION  DESCRIPTION
    --------------------------------------------------------------------
    2026-10-21  미지정      v0.1.0   파일 역할·입출력·가정·이력 명세 추가

    2026-10-21  김재현(Finn) v0.1.0   Agent 도구 선택 기반 리빌드·업무 경계 및 검증

    2026-10-22  김재현(Finn) v0.1.0   자료·Agent·검수 승인·식품 DB 중심 흐름 및 알레르기 표기 정리

    상세 Git 이력: git log --follow -- backend/tests/test_agent_routing.py
    공통 규칙: doc/file-header-policy.md

기존 모듈 설명:
    Agent decisions must change actual tools, preserve evidence, and stay bounded.
==============================================================================
"""

# FMS_MODULE_METADATA_BEGIN
__author__ = "김재현(Finn)"
__email__ = "penguin.klg@gmail.com"
__version__ = "0.1.0"
__status__ = "Development"
# FMS_MODULE_METADATA_END

# ---------------------------------------------------------------------------
# 1. Imports / Module implementation
# ---------------------------------------------------------------------------

import asyncio
from copy import deepcopy
import pytest
from backend.app.infrastructure.database import Problem
from backend.app.domains.analysis.planning import sources,decide,selected_sources,image_regions,group_references,restore_groups
from backend.app.domains.analysis.agent import batches


def sample(kind='html'):
    return [{'text':'첫 상품 100g 2000원','evidence_id':'E1','asset_id':'A','locator':{'kind':kind,'page_number':1,'bbox':[0,.1,.4,.2]}},
            {'text':'둘째 상품 200g 3000원','evidence_id':'E2','asset_id':'A','locator':{'kind':kind,'page_number':1,'bbox':[.6,.1,1,.2]}}]


def test_decision_selects_real_tool_and_invalid_missing_or_duplicate_sources_fail():
    rows=sample();rows[1]['asset_id']='B'
    async def model(task,payload,schema):
        return {'document_type':'mixed','summary':'자료별 판단','tasks':['match_catalog'],'routes':[
            {'source_id':'D1','tool':'extract_text','reason':'본문 정보'},
            {'source_id':'D2','tool':'request_review','reason':'자료 불충분'}]}
    plan=asyncio.run(decide(rows,'공지',model))
    assert [r['tool'] for r,s in selected_sources(plan,rows)]==['extract_text','request_review']
    for altered in [plan['routes'][:1],[plan['routes'][0],plan['routes'][0]]]:
        bad={**plan,'routes':altered}
        with pytest.raises(Problem) as error:selected_sources(bad,rows)
        assert error.value.code=='AGENT_PLAN_COVERAGE'
    bad=deepcopy(plan);bad['routes'][0]['tool']='normalize_table'
    with pytest.raises(Problem) as error:selected_sources(bad,rows)
    assert error.value.code=='AGENT_TOOL_UNAVAILABLE'


def test_image_layout_uses_coordinates_and_keeps_products_separate():
    source=sources(sample('image'))[0]
    async def model(task,payload,schema):
        assert payload['lines'][0][1]==[0,100,400,200]
        return {'regions':[{'label':'첫 상품','source_ids':['1'],'context_ids':[]},{'label':'둘째 상품','source_ids':['2'],'context_ids':[]}],'unassigned_ids':[]}
    groups,manual,layout=asyncio.run(image_regions(source,model))
    assert not manual and [[s['evidence_id'] for s in g] for g in groups]==[['E1'],['E2']]
    assert groups[0][0]['region_label']=='첫 상품'


def test_image_layout_cannot_silently_omit_source_or_merge_duplicate_regions():
    async def model(task,payload,schema):return {'regions':[{'label':'상품','source_ids':['1'],'context_ids':[]}],'unassigned_ids':[]}
    with pytest.raises(Problem) as error:asyncio.run(image_regions(sources(sample('image'))[0],model))
    assert error.value.code=='IMAGE_REGION_COVERAGE'


def test_window_references_resume_without_copying_source_text_into_checkpoint():
    rows=sample();rows[0]['text']='식품'*3000
    groups=list(batches(rows));refs=group_references(groups)
    assert all('text' not in item for group in refs for item in group)
    restored=restore_groups(refs,rows)
    assert ''.join(s['text'] for g in restored for s in g)==''.join(s['text'] for s in rows)


def test_compact_product_region_preserves_all_ocr_evidence_without_per_line_overhead():
    from backend.app.domains.analysis.agent import model_mapping,extraction_payload
    from backend.app.domains.analysis.provider import LocalModel
    from backend.app.domains.analysis.schemas import event_schema
    rows=[{'text':'상품 원재료 규격 정보 '+str(i),'evidence_id':'E'+str(i),'region_label':'하나의 상품','locator':{'kind':'image','page_number':1,'bbox':[0,0,1,1]}} for i in range(41)]
    mapping=model_mapping(rows)
    assert list(mapping)==['S1']
    assert mapping['S1']['text']=='\n'.join(row['text'] for row in rows)
    assert [mapping['S1']['evidence_id'],*mapping['S1']['related_evidence_ids']]==[row['evidence_id'] for row in rows]
    assert len(extraction_payload({'title':'상품','assets':[]},'marketing',mapping)['segments'])==1
    LocalModel().request('extract',extraction_payload({'title':'상품','assets':[]},'marketing',mapping),event_schema(list(mapping)))


def test_redundant_unassigned_ids_keep_valid_region_assignment():
    async def model(*args):return {'regions':[{'label':'상품','source_ids':['1','2'],'context_ids':[]}],'unassigned_ids':['1','2']}
    groups,manual,layout=asyncio.run(image_regions(sources(sample('image'))[0],model))
    assert len(groups)==1 and not manual and layout['unassigned_ids']==[]
    assert [s['evidence_id'] for s in groups[0]]==['E1','E2']


def test_job_stage_log_is_bounded_and_does_not_repeat_heartbeat_updates():
    from backend.app.domains.jobs.service import stamp_stage
    job={'stage':'수집 시작'}
    stamp_stage(job);stamp_stage(job)
    assert len(job['stage_history'])==1
    for i in range(150):job['stage']='처리 '+str(i);stamp_stage(job)
    assert len(job['stage_history'])==100
    assert job['stage_history'][0]['stage']=='처리 50' and job['stage_history'][-1]['stage']=='처리 149'


def test_merged_cell_memory_budget_is_shared_across_sheets(tmp_path,monkeypatch):
    import openpyxl
    from backend.app.parsers import xlsx_merges
    book=openpyxl.Workbook();book.active.merge_cells('A1:B3');book.create_sheet('두번째').merge_cells('A1:B3')
    path=tmp_path/'merged.xlsx';book.save(path);book.close()
    monkeypatch.setattr(xlsx_merges,'MAX_MERGE_CELLS',10)
    with pytest.raises(Problem) as error:xlsx_merges.merge_anchors(path)
    assert error.value.code=='SHEET_LIMIT'
