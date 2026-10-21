"""
==============================================================================
FMS_FILE_SPEC
export.py - 근거를 포함한 검토 Excel 생성
==============================================================================
설명:
    - 근거를 포함한 검토 Excel 생성
입력·출력 명세:
    - 입력: 분석 보고서·검토된 식품 후보·원문 근거
    - 출력: Excel bytes 또는 저장한 export asset 참조
실행 가정:
    - 원문 문자열을 수식으로 실행하지 않고 DB 승인과 구분
작성 정보:
    - 작성자: 미지정
    - 이메일: 미지정
    - 작성일: 2026-10-21 (최초 Git 등록일 기준)
    - 라이선스: 미지정
    - 버전: 0.1.0
    - 상태: Development
변경 이력 (Revision History):
    DATE        AUTHOR      VERSION  DESCRIPTION
    --------------------------------------------------------------------
    2026-10-21  미지정      v0.1.0   최초 Git 등록 a11ef05a
    2026-10-21  미지정      v0.1.0   파일 역할·입출력·가정·이력 명세 추가

    상세 Git 이력: git log --follow -- backend/app/domains/analysis/export.py
    공통 규칙: doc/file-header-policy.md

기존 모듈 설명:
    Automatic, evidence-linked review workbook; source strings are never formulas.
==============================================================================
"""

# FMS_MODULE_METADATA_BEGIN
__author__ = "미지정"
__email__ = "미지정"
__version__ = "0.1.0"
__status__ = "Development"
# FMS_MODULE_METADATA_END

# ---------------------------------------------------------------------------
# 1. Imports / Module implementation
# ---------------------------------------------------------------------------

from io import BytesIO
from contextlib import nullcontext
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment
from backend.app.infrastructure.database import get
from backend.app.infrastructure.database import now
from backend.app.infrastructure.database import put
from backend.app.infrastructure.database import tx
from backend.app.domains.ingestion.collection import store_asset

def workbook_bytes(report,evidence,source_url):
    wb=Workbook();ws=wb.active;ws.title='식품 변경안'
    labels={'food_profile':'식품 기준정보','price':'단가','temporarily_unavailable':'공급 일시중단','supply_ended':'공급 종료','discontinued':'단종','available':'공급 재개','substitution':'대체품','new_product':'신제품'}
    ws.append(['변경 종류','식품명','규격','금액(원)','기준 수량','단위','적용 시작일','종료일','사유','대체품 후보','가격 열 기준','DB 대조','검토 필요 사항','변경안 ID','근거 ID','원문 링크','검토 상태','검토 메모'])
    for original in report['events']:
        from backend.app.domains.catalog.record import upgrade_event
        e=upgrade_event({**original,**original.get('review',{}).get('fields',{})})
        ws.append([labels[e['kind']],e['product_name'],e['spec'],e['amount'],e['quantity'],e['unit'],e['effective_from'],e['effective_to'],e['reason'],' / '.join(e['replacements']),e['price_basis'],e.get('match_status'), '\n'.join(e.get('review_reasons',[])),e.get('proposal_id'),'\n'.join(e['evidence_ids']),source_url,e.get('review',{}).get('status','pending'),e.get('review',{}).get('reason')])
    from backend.app.domains.catalog.record import from_event; from backend.app.domains.catalog.record import ALLERGENS
    foods=wb.create_sheet('표준 식품 정보');foods.append(['공통식품코드','식품명','원문 식품명','업체','맛·유형','개별 수량','개별 단위','포장 개수','총 수량','총 단위','환산 수량','환산 단위','판매 단위','크기','품질 등급','용도','표시 기준','원재료','원산지','알레르기 함유 표시','같은 시설·혼입 표시','알레르기 후보','알레르기 확인 상태','자료 적용가','조사가1','조사가2','조사가3','비고'])
    def labels_for(codes):return ', '.join(ALLERGENS[c][0] for c in codes) if codes is not None else None
    for event in report['events']:
        food=event.get('review',{}).get('fields',{}).get('food_record') or from_event(event);a=food['attributes'];s=food['specification'];al=food['allergens'];prices=food['prices']
        foods.append([food['food_code'],food['food_name'],food['raw_food_name'],food['company_name'],a['flavor'],s['piece_quantity'],s['piece_unit'],s['pack_count'],s['quantity'],s['unit'],s['normalized_quantity'],s['normalized_unit'],s['sales_unit'],a['size_grade'],a['quality_grade'],a['intended_use'],a['labeling_requirement'],__import__('json').dumps(a['ingredients'],ensure_ascii=False) if a['ingredients'] else None,a['origin_country'],labels_for(al['contains']),labels_for(al['may_contain']),labels_for(al['candidates']),al['review_status'],prices['source_applied'],prices['survey_1'],prices['survey_2'],prices['survey_3'],food['note']])
    wb.move_sheet(foods,offset=-1);wb.active=0
    evs=wb.create_sheet('원문 근거');evs.append(['근거 ID','원문 인용','위치·시트·셀·페이지','OCR 신뢰도','원문 링크'])
    for e in evidence:evs.append([e['id'],e['quote'],__import__('json').dumps(e['locator'],ensure_ascii=False),e.get('ocr_confidence'),e.get('source_url')])
    info=wb.create_sheet('AI 실행 정보');info.append(['항목','값']);info.append(['모델',report['model']]);info.append(['AI 처리 요약',report['plan']['summary']]);info.append(['원문 구간 수',report['coverage']['source_segments']]);info.append(['처리 구간 수',report['coverage']['processed_segments']]);info.append(['실제 모델 호출 수',len(report['model_calls'])]);info.append(['안내','검토용 추출 결과입니다. 승인·반영 전까지 DB 확정 단가가 아닙니다.']);info.append(['분석 실행 ID',report['analysis_run_id']])
    for call in report['model_calls']:info.append([call['task'],f"{call['model']} · 입력 {call['input_tokens']} / 출력 {call['output_tokens']} 토큰 · {call['elapsed_seconds']}초"])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='234536');cell.font=Font(color='FFFFFF',bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):cell.value=cell.value[:32767];cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for col in sheet.columns:
            letter=col[0].column_letter;sheet.column_dimensions[letter].width=min(54,max(16,max(len(str(c.value or '')) for c in list(col)[:30])+2))
    stream=BytesIO();wb.save(stream);return stream.getvalue()

def export_review(state,connection=None):
    from backend.app.domains.analysis.agent import progress
    if not state.get('review_mode'):progress(state['run_id'],'검토용 엑셀 자동 생성 · 원문 근거와 AI 실행 내역 연결')
    with (nullcontext(connection) if connection is not None else tx()) as c:
        report=get(c,'agent_report',state['agent_report_id'],True)
        if report.get('export_asset_id') and report.get('export_dto_schema')=='food-record-v2':return {'export_asset_id':report['export_asset_id']}
        if report.get('export_asset_id'):report.setdefault('previous_export_asset_ids',[]).append(report['export_asset_id'])
        evidence_ids=list(dict.fromkeys(eid for event in report['events'] for eid in event['evidence_ids']))
        evidence=[get(c,'evidence',eid) for eid in evidence_ids];r=get(c,'revision',state['revision_id']);doc=get(c,'document',r['document_id'])
        content=workbook_bytes(report,evidence,doc.get('source_url'))
        asset=store_asset(c,content,'AI_식품정보_검토_'+state['run_id'][:8]+'.xlsx')
        report.update(export_asset_id=asset['id'],export_dto_schema='food-record-v2',status='completed',finished_at=now());put(c,'agent_report',report)
        run=get(c,'analysis_run',state['run_id'],True);run['export_asset_id']=asset['id'];put(c,'analysis_run',run)
    return {'export_asset_id':asset['id']}
