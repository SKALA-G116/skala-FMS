"""Source workbook reader. Never writes authoritative product/price facts.
Read-only streaming import; values, sheet/row/cell coordinates and formula absence
are retained. Queries, compound filters, ordering and pagination run in PostgreSQL.
"""
import json
import re
import zipfile
from contextlib import ExitStack
from decimal import Decimal
from decimal import InvalidOperation
from datetime import date
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter
from psycopg.types.json import Jsonb
from pathlib import Path
from backend.app.infrastructure.database import Problem
from backend.app.infrastructure.database import get
from backend.app.infrastructure.database import uid
from backend.app.domains.ingestion.documents import DATA
from backend.app.domains.ingestion.documents import import_document

MAX_ROWS=20000

MAX_CELLS=200000

MAX_COLS=100

MAX_SHEETS=20

def metadata(c,wid):
    w=c.execute('SELECT * FROM price_workbooks WHERE id=%s',(wid,)).fetchone()
    if not w:raise Problem(404,'NOT_FOUND','단가표 파일이 없습니다.')
    sheets=c.execute('SELECT * FROM price_sheets WHERE workbook_id=%s ORDER BY position',(wid,)).fetchall()
    return {'id':str(w['id']),'document_id':str(w['document_id']),'document_revision_id':str(w['revision_id']),'asset_id':str(w['asset_id']) if w['asset_id'] else None,'filename':w['filename'],'content_sha256':w['content_sha256'],'created_at':w['created_at'].isoformat(),'sheets':[{'id':str(s['id']),'name':s['name'],'header_row':s['header_row'],'row_count':s['row_count'],'columns':s['columns']} for s in sheets]}

def register(c,revision_id,asset_id=None):
    rev=get(c,'revision',revision_id);asset_id=asset_id or rev['snapshot_asset_id']
    if asset_id not in [a['id'] for a in rev['assets']]:raise Problem(422,'ASSET_MISMATCH','이 자료에 속하는 첨부를 선택하세요.')
    old=c.execute('SELECT id FROM price_workbooks WHERE revision_id=%s AND asset_id=%s',(revision_id,asset_id)).fetchone()
    if old:return metadata(c,str(old['id']))
    a=get(c,'asset',asset_id)
    if Path(a['filename']).suffix.lower()!='.xlsx':raise Problem(422,'WORKBOOK_FORMAT','원본 표 보기는 XLSX 파일을 지원합니다.')
    path=DATA/a['storage_key'];streams=ExitStack()
    try:
        with zipfile.ZipFile(path) as z:
            if len(z.infolist())>5000 or sum(i.file_size for i in z.infolist())>100*1024*1024:raise Problem(422,'EXPANDED_LIMIT','압축 해제 크기가 너무 큽니다.')
        book=openpyxl.load_workbook(streams.enter_context(path.open('rb')),read_only=True,data_only=True,keep_links=False)
        formulas=openpyxl.load_workbook(streams.enter_context(path.open('rb')),read_only=True,data_only=False,keep_links=False)
    except Exception as e:
        streams.close()
        if isinstance(e,Problem):raise
        raise Problem(422,'INVALID_WORKBOOK','읽을 수 없는 XLSX 파일입니다.') from e
    wid=uid();nonempty=0
    c.execute('INSERT INTO price_workbooks(id,document_id,revision_id,filename,content_sha256,asset_id) VALUES(%s,%s,%s,%s,%s,%s)',(wid,rev['document_id'],revision_id,a['filename'],a['sha256'],asset_id))
    try:
        if len(book.worksheets)>MAX_SHEETS:raise Problem(422,'SHEET_LIMIT','시트는 20개까지 지원합니다.')
        for pos,sheet in enumerate(book):
            if sheet.max_column>MAX_COLS or sheet.max_row>MAX_ROWS+1:raise Problem(422,'SHEET_LIMIT','시트당 2만 행·100열까지 지원합니다.')
            iterator=sheet.iter_rows();header=next(iterator,None)
            formula_rows=formulas[sheet.title].iter_rows();next(formula_rows,None)
            if header is None:continue
            columns=[]
            for i,cell in enumerate(header):
                letter=get_column_letter(i+1);raw='' if cell.value is None else str(cell.value)
                label=re.sub(r'\s+',' ',raw.replace('_x000D_',' ').replace('_x000a_',' ')).strip()
                columns.append({'key':letter,'label':label or f'제목 없는 열 ({letter})','original_label':raw,'numeric':False,'format':cell.number_format})
            sid=uid();row_count=0;numeric_counts=[0]*len(columns);text_counts=[0]*len(columns)
            c.execute('INSERT INTO price_sheets(id,workbook_id,name,position,header_row,columns,row_count) VALUES(%s,%s,%s,%s,%s,%s,%s)',(sid,wid,sheet.title,pos,1,Jsonb(columns),0))
            batch=[]
            for rownum,row in enumerate(iterator,2):
                formula_row=next(formula_rows,())
                values={};numbers={}
                for i,cell in enumerate(row):
                    value=cell.value
                    if value is None and i<len(formula_row) and formula_row[i].data_type=='f':value='수식 (저장된 결과 없음): '+str(formula_row[i].value)
                    if value is None:continue
                    nonempty+=1
                    if nonempty>MAX_CELLS:raise Problem(422,'CELL_LIMIT','파일의 비어 있지 않은 셀이 20만 개를 초과했습니다.')
                    letter=columns[i]['key']
                    if isinstance(value,(datetime,date)):text=value.isoformat()
                    elif isinstance(value,bool):text='TRUE' if value else 'FALSE'
                    else:text=str(value)
                    if len(text)>32768:raise Problem(422,'CELL_TEXT_LIMIT','한 셀의 내용이 너무 깁니다.')
                    values[letter]=text
                    if isinstance(value,(int,float,Decimal)) and not isinstance(value,bool):
                        n=Decimal(str(value))
                        if n.is_finite():numbers[letter]=str(n);numeric_counts[i]+=1
                    else:text_counts[i]+=1
                if not values:continue
                batch.append((sid,rownum,Jsonb(values),Jsonb(numbers)));row_count+=1
                if len(batch)==200:
                    with c.cursor() as cur:cur.executemany('INSERT INTO price_sheet_rows VALUES(%s,%s,%s,%s)',batch)
                    batch=[]
            if batch:
                with c.cursor() as cur:cur.executemany('INSERT INTO price_sheet_rows VALUES(%s,%s,%s,%s)',batch)
            for i,column in enumerate(columns):column['numeric']=numeric_counts[i]>0 and text_counts[i]==0
            c.execute('UPDATE price_sheets SET columns=%s,row_count=%s,source_header=%s WHERE id=%s',(Jsonb(columns),row_count,Jsonb(columns),sid));detect_header(c,sid)
    finally:book.close();formulas.close();streams.close()
    return metadata(c,wid)

def import_workbook(c,title,content,filename):
    result=import_document(c,title,content,filename)
    return register(c,result['revision']['id'])

def query_rows(c,sheet_id,query):
    sheet=c.execute('SELECT * FROM price_sheets WHERE id=%s',(sheet_id,)).fetchone()
    if not sheet:raise Problem(404,'NOT_FOUND','시트가 없습니다.')
    columns={col['key']:col for col in sheet['columns']}
    try:
        limit=int(query.get('limit',50));offset=int(query.get('offset',0));filters=json.loads(query.get('filters','[]'))
        if not 1<=limit<=100 or offset<0 or not isinstance(filters,list) or len(filters)>20:raise ValueError()
    except (ValueError,TypeError):raise Problem(422,'INVALID_FILTER','필터나 페이지 범위를 확인하세요.')
    where=['sheet_id=%s','row_number>%s'];args=[sheet_id,sheet['header_row']]
    for f in filters:
        if not isinstance(f,dict) or f.get('column') not in columns:raise Problem(422,'INVALID_COLUMN','필터 열을 확인하세요.')
        column=f['column'];op=f.get('op','contains');value=f.get('value','')
        if not isinstance(value,str) or len(value)>500:raise Problem(422,'INVALID_FILTER','필터 값은 500자까지 입력할 수 있습니다.')
        expr=f"cells->>'{column}'" # column is validated against generated A..CV keys
        if op in ('empty','not_empty'):where.append(f"COALESCE({expr},'') {'=' if op=='empty' else '<>'} ''")
        elif op in ('contains','equals'):
            if op=='contains':
                where.append(f"strpos(lower(COALESCE({expr},'')),lower(%s))>0");args.append(value)
            else:where.append(f"{expr}=%s");args.append(value)
        elif op in ('gte','lte','eq_number'):
            try:n=Decimal(value);assert n.is_finite() and len(value)<=50
            except (InvalidOperation,AssertionError):raise Problem(422,'INVALID_NUMBER','유효한 숫자를 입력하세요.')
            if not columns[column]['numeric']:raise Problem(422,'NOT_NUMERIC','숫자 범위는 숫자 열에만 사용할 수 있습니다.')
            operator={'gte':'>=','lte':'<=','eq_number':'='}[op];where.append(f"(numbers->>'{column}')::numeric {operator} %s::numeric");args.append(str(n))
        else:raise Problem(422,'INVALID_OPERATOR','지원하지 않는 필터 조건입니다.')
    search=query.get('q','').strip()
    if len(search)>500:raise Problem(422,'INVALID_FILTER','검색어는 500자까지 입력할 수 있습니다.')
    if search:where.append("EXISTS (SELECT 1 FROM jsonb_each_text(cells) v WHERE strpos(lower(v.value),lower(%s))>0)");args.append(search)
    order=query.get('sort','row_number');direction=query.get('direction','asc')
    if direction not in ('asc','desc') or (order!='row_number' and order not in columns):raise Problem(422,'INVALID_SORT','정렬 열을 확인하세요.')
    if order=='row_number':sort=f'row_number {direction}'
    else:
        expr=f"(numbers->>'{order}')::numeric" if columns[order]['numeric'] else f"cells->>'{order}'"
        sort=f'{expr} {direction} NULLS LAST,row_number ASC'
    sqlwhere=' AND '.join(where)
    count=c.execute('SELECT count(*) n FROM price_sheet_rows WHERE '+sqlwhere,args).fetchone()['n']
    rows=c.execute('SELECT row_number,cells FROM price_sheet_rows WHERE '+sqlwhere+' ORDER BY '+sort+' LIMIT %s OFFSET %s',[*args,limit,offset]).fetchall()
    return {'sheet_id':str(sheet['id']),'items':[{'row_number':r['row_number'],'cells':r['cells'],'cell_range':f"A{r['row_number']}:{sheet['columns'][-1]['key']}{r['row_number']}"} for r in rows],'total':count,'sheet_row_count':sheet['row_count'],'offset':offset,'limit':limit,'has_more':offset+limit<count}

def set_header(c,sheet_id,header_row):
    if not isinstance(header_row,int) or not 1<=header_row<=20000:raise Problem(422,'HEADER_ROW','제목 행 번호를 확인하세요.')
    s=c.execute('SELECT * FROM price_sheets WHERE id=%s FOR UPDATE',(sheet_id,)).fetchone()
    if not s:raise Problem(404,'NOT_FOUND','시트가 없습니다.')
    original=s['source_header'] or s['columns']
    if header_row==1:values={col['key']:col['original_label'] for col in original}
    else:
        row=c.execute('SELECT cells FROM price_sheet_rows WHERE sheet_id=%s AND row_number=%s',(sheet_id,header_row)).fetchone()
        if not row:raise Problem(422,'HEADER_ROW','내용이 있는 제목 행을 지정하세요.')
        values=row['cells']
    stats={r['key']:r for r in c.execute("SELECT v.key,count(*) total,count(numbers->>v.key) numeric_count FROM price_sheet_rows CROSS JOIN LATERAL jsonb_each_text(cells) v WHERE sheet_id=%s AND row_number>%s GROUP BY v.key",(sheet_id,header_row))}
    columns=[]
    for col in original:
        raw=values.get(col['key'],'');label=re.sub(r'\s+',' ',raw.replace('_x000D_',' ')).strip();st=stats.get(col['key'],{})
        columns.append({**col,'label':label or f"제목 없는 열 ({col['key']})",'original_label':raw,'numeric':bool(st.get('numeric_count') and st['numeric_count']==st['total'])})
    count=c.execute('SELECT count(*) n FROM price_sheet_rows WHERE sheet_id=%s AND row_number>%s',(sheet_id,header_row)).fetchone()['n']
    c.execute('UPDATE price_sheets SET source_header=COALESCE(source_header,columns),columns=%s,header_row=%s,row_count=%s WHERE id=%s',(Jsonb(columns),header_row,count,sheet_id))
    return metadata(c,str(s['workbook_id']))

def detect_header(c,sheet_id):
    s=c.execute('SELECT columns,source_header FROM price_sheets WHERE id=%s',(sheet_id,)).fetchone()
    candidates=[(1,{col['key']:col['original_label'] for col in (s['source_header'] or s['columns'])})]+[(r['row_number'],r['cells']) for r in c.execute('SELECT row_number,cells FROM price_sheet_rows WHERE sheet_id=%s AND row_number<=20 ORDER BY row_number',(sheet_id,))]
    def score(values):return sum(bool(re.search(r'식품명|상품명|제품명|품명|규격|단위|학교가|판매가|단가|적용가|조사가|업체군',v)) for v in values.values())
    best=max(candidates,key=lambda pair:score(pair[1]))
    if best[0]!=1 and score(best[1])>=3:return set_header(c,sheet_id,best[0])
    return None
